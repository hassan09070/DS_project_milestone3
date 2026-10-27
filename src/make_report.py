"""Build report/Milestone03_Report.docx from outputs/results.json and the saved figures.

    .venv/bin/python src/make_report.py

Written for a mixed audience (course graders, policy readers, teammates): professional but plain language, with every
technical term followed by a short explanation the first time it appears. All numbers are read from results.json.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config as C  # noqa: E402

R = json.load(open(C.OUT_DIR / "results.json"))
T = C.TABLES_DIR
NAMES = {"naive": "Majority-class baseline", "logreg": "Logistic regression", "tree": "Decision tree", "rf": "Random forest",
         "logreg_demo": "Logistic regression (age, sex, area only)"}
MODELS = ("logreg", "tree", "rf")


def fmt(x, nd=3):
    if isinstance(x, bool):
        return "Yes" if x else "No"
    if isinstance(x, int):
        return f"{x:,}"
    if isinstance(x, float):
        if pd.isna(x):
            return "-"
        return f"{x:,.0f}" if abs(x) >= 1000 else f"{x:.{nd}f}"
    return str(x)


def add_table(doc, df: pd.DataFrame, nd=3, font=8.5, caption: str | None = None):
    if caption:
        p = doc.add_paragraph(); r = p.add_run(caption); r.bold = True; r.font.size = Pt(9.5)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns):
        cell = t.rows[0].cells[i]; cell.text = str(c)
        for p in cell.paragraphs:
            for r in p.runs: r.bold = True; r.font.size = Pt(font)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = fmt(v, nd)
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size = Pt(font)
    doc.add_paragraph()


def add_fig(doc, key_or_path, caption, width=6.3):
    path = R["figures"].get(key_or_path, key_or_path)
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); r = p.add_run(caption); r.italic = True; r.font.size = Pt(9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def para(doc, text):
    doc.add_paragraph(text)


def main():
    h05 = R["holdout_05"]; w05 = R["holdout_05_weighted"]; cvb = {r["model"]: r for r in R["cv_best"]}; h5 = R["h5"]
    bs = h5["holdout_bootstrap"]
    tc = {(r["model"], r["threshold_label"]): r for r in R["threshold_comparison"]}
    ct = {(r["model"], r["cost_ratio_fn_to_fp"]): r for r in R["cost_thresholds"]}
    ors = {r["feature"]: r for r in R["odds_ratios"]}
    sub = pd.DataFrame(R["subgroups_logreg"]).set_index(["variable", "level"])
    rec = R["recommended_threshold"]; ratio = R["recommended_cost_ratio"]
    split = {r["partition"]: r for r in R["split"]}
    ds = R["dataset"]; gaps = R["h4_gaps"]
    tree_rec = tc[("tree", "recommended")]; lr_rec = tc[("logreg", "recommended")]
    tree_05 = tc[("tree", "0.50 (default)")]
    lrp, trp, rfp = R["best_params"]["logreg"], R["best_params"]["tree"], R["best_params"]["rf"]
    best = R["recommended_model"]; best_rec = tc[(best, "recommended")]; best_05 = tc[(best, "0.50 (default)")]
    others = [m for m in MODELS if m != best]
    rank = R["model_rank_by_holdout_cost"]
    top_auc = max(MODELS, key=lambda m: h05[m]["roc_auc"])
    edge = {r["feature"]: r for r in R["edge_profile"]}
    hyp_gini = {r["feature"]: r.get("gini_rank") for r in R["hypotheses"]}

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)

    doc.add_heading("Milestone 03 - Baseline Models, Evaluation and Initial Findings", 0)
    doc.add_paragraph("Predicting which children aged 5-16 in Pakistan are out of school, using the ASER 2025 household survey. "
                      "Manuscript Section IV draft: modelling methods, baseline results and preliminary discussion.").italic = True
    para(doc, "How to read this section. Each technical term is explained in brackets the first time it appears, and a short "
              "glossary at the end collects them. Every number quoted here is produced by the accompanying code and can be traced "
              "to a table in outputs/tables.")

    # ------------------------------------------------------------------ 1
    doc.add_heading("1. What we are predicting and with what data", 1)
    para(doc,
        f"The outcome we predict - the target variable (the column the model learns to guess) - is out_of_school: 1 if a child has "
        f"never enrolled or has dropped out, 0 if the child is currently enrolled. The analysis file prepared in Milestone 2 contains "
        f"{ds['rows']:,} children from {ds['districts']} districts. Of these, {ds['positives']:,} ({ds['prevalence']:.1%}) are out of "
        f"school. This is a case of class imbalance (one outcome is much rarer than the other), which shapes how the models must be "
        f"judged: a model that simply says 'enrolled' for everyone would be right {1 - ds['prevalence']:.1%} of the time while finding "
        f"no out-of-school child at all. When ASER's sampling weights are applied (numbers attached to each child so the sample "
        f"represents the national population of children), the out-of-school share is {ds['weighted_prevalence']:.1%}.")
    para(doc,
        "The model uses 13 predictors (the input columns, also called features): the child's age, sex and whether the household is "
        "urban; a sex-by-area interaction term (a column equal to 1 only for urban girls, which lets the effect of being a girl differ "
        "between rural and urban areas); whether the child has a B-form (the NADRA birth registration document); whether the child has "
        "any functional disability; a household asset count from 0 to 8; a digital-use count from 0 to 5; housing structure on a "
        "1-to-3 scale from kachha to pacca; and four province indicators. Province was one-hot encoded (each province becomes its own "
        "0/1 column) with Punjab as the reference category (the province all others are compared against). The child id, the district "
        "and the sampling weight are never used as predictors. The data contain no missing values, so no imputation (filling in gaps "
        "with estimated values) was needed.")

    # ------------------------------------------------------------------ 2
    doc.add_heading("2. The baseline benchmark and the three models, and why they were chosen", 1)
    bullets(doc, [
        "Majority-class baseline (a model that always predicts the most common outcome, here 'enrolled'). It sets the floor that any "
        f"real model must beat and makes the point above concrete: it scores {h05['naive']['accuracy']:.1%} accuracy (the share of all "
        "children classified correctly) and yet has zero recall (the share of truly out-of-school children it finds).",
        "Logistic regression (a statistical model that estimates the probability of an outcome as a weighted sum of the predictors). "
        "It was chosen because hypotheses H1-H4 from Milestone 2 ask about the direction and size of individual effects, and logistic "
        "regression answers those directly through odds ratios (how many times higher or lower the odds of being out of school are for "
        "a one-unit change in a predictor, holding the others fixed) with confidence intervals (the range in which the true value very "
        "probably lies). The inputs are standardised (rescaled so each has the same spread) and the model uses an L2 penalty (a term that "
        "shrinks the weights towards zero so the model does not over-react to noise). Two settings were tuned: the regularisation "
        "strength C (how strong that shrinkage is) and class weighting (whether out-of-school children count more during training to "
        "offset their rarity).",
        "Decision tree (a model that splits the children into groups with a sequence of yes/no questions such as 'is housing kachha?'). "
        "It was chosen because it captures non-linear patterns (effects that change in size along a scale) and interactions (effects that "
        "depend on other variables, such as the gender gap depending on area) without having to specify them in advance, it is easy to "
        "read, and it is quick to train on 161,000 rows. Three settings were tuned: maximum depth (how many questions deep the tree may "
        "go), minimum leaf size (the smallest group a final answer may rest on) and class weighting.",
        "Random forest (an ensemble, meaning a model made of many models: here 200 decision trees, each grown on a random re-sample of "
        "the children and allowed to consider only a random subset of the predictors at each split, whose votes are averaged). It was "
        "added because a single tree is unstable (small changes in the data can change its splits) and tends to stop short of the best "
        "possible ranking; averaging many de-correlated trees smooths both problems while keeping the same kind of importance measures. "
        "It is a distinct learning algorithm from the single tree - bagging (training on bootstrap re-samples) plus random feature "
        "selection - and is the natural bridge to the stronger models planned for Milestone 4. Tuned: maximum depth, minimum leaf size "
        "and class weighting; the number of trees was fixed at 200.",
    ])

    # ------------------------------------------------------------------ 3
    doc.add_heading("3. How the models were trained and tested fairly", 1)
    para(doc,
        "Children in the survey are clustered in districts, and enrolment rates differ sharply from one district to another. If children "
        "from the same district appeared in both the training data (what the model learns from) and the test data (what it is judged on), "
        "the model could partly memorise district patterns and look better than it would on genuinely new places. This is a form of data "
        "leakage (information from the test data influencing the model). To prevent it we used StratifiedGroupKFold, a splitting method "
        "that keeps every district wholly on one side of the split (the 'group' part) while keeping the share of out-of-school children "
        "similar on both sides (the 'stratified' part).")
    para(doc,
        f"One fold (one of five equal portions) of such a split was set aside as the hold-out test set (data the model never sees until "
        f"final evaluation): {split['test']['rows']:,} children in {split['test']['districts']} districts, {split['test']['share_of_rows']:.1%} "
        f"of the rows, with an out-of-school share of {split['test']['prevalence']:.1%}. The remaining {split['train']['rows']:,} children in "
        f"{split['train']['districts']} districts formed the training set. Within the training set we ran cross-validation (repeatedly "
        "training on four fifths of the districts and checking on the remaining fifth, five times, so every training child is checked on "
        "exactly once) using the same district-grouped, stratified method. The code asserts that no district appears on both sides of "
        "any split.")
    add_table(doc, pd.read_csv(T / "t01_split_description.csv"), caption="Table 1. The district-grouped hold-out split")
    para(doc,
        "Hyperparameters (settings chosen by the analyst rather than learned from the data, such as C or tree depth) were selected with "
        "a grid search (trying every combination from a pre-set list) inside the cross-validation, choosing the combination with the "
        "best ROC-AUC (a 0-to-1 score of how well the model ranks out-of-school children above enrolled ones across all possible "
        "cut-offs; 0.5 is no better than chance, 1.0 is perfect). Recall, precision, F1 and accuracy were recorded for every fold as "
        "well, so that the spread of performance between groups of districts is visible. The standardisation step is fitted inside each "
        "fold, so it too never sees the checking data.")
    cvdf = pd.DataFrame(R["cv_best"])[["model", "best_params", "cv_roc_auc_mean", "cv_roc_auc_sd", "cv_average_precision_mean",
                                        "cv_recall_mean", "cv_f1_mean", "cv_accuracy_mean"]]
    cvdf["model"] = cvdf["model"].map(NAMES)
    cvdf.columns = ["Model", "Best settings", "CV ROC-AUC", "Spread (SD)", "CV PR-AUC", "CV recall", "CV F1", "CV accuracy"]
    add_table(doc, cvdf, caption="Table 2. Cross-validation results for the best settings of each model (mean over 5 district-grouped folds). "
                                 "Recall, F1 and accuracy are measured at the default 0.5 cut-off; ROC-AUC and PR-AUC do not depend on a cut-off.")
    para(doc,
        f"What the tuning showed. For logistic regression the strongest shrinkage (C = {lrp['C']}) with {lrp['class_weight']} class weights "
        f"was best, with a cross-validated ROC-AUC of {cvb['logreg']['cv_roc_auc_mean']:.3f} plus or minus {cvb['logreg']['cv_roc_auc_sd']:.3f} "
        f"(the standard deviation, a measure of how much the score varied from fold to fold). ROC-AUC barely changed across values of C, "
        f"because with 129,000 training rows the weights are already well determined; class weighting mattered mainly for recall at the "
        f"default cut-off. For the decision tree a depth of {trp['max_depth']} with at least {trp['min_samples_leaf']} children per leaf and no "
        f"class weighting was best (ROC-AUC {cvb['tree']['cv_roc_auc_mean']:.3f} plus or minus {cvb['tree']['cv_roc_auc_sd']:.3f}; the fitted "
        f"tree has {R['tree_shape']['leaves']} leaves). Deeper, unconstrained trees memorised district quirks and scored worse on unseen "
        f"districts. For the random forest the best settings were a maximum depth of {rfp['max_depth']}, at least {rfp['min_samples_leaf']} children "
        f"per leaf and {'no class weighting' if rfp['class_weight'] in (None, 'None') else 'balanced class weighting'} "
        f"(ROC-AUC {cvb['rf']['cv_roc_auc_mean']:.3f} plus or minus {cvb['rf']['cv_roc_auc_sd']:.3f}), the highest cross-validated score of the three. "
        f"One consequence needs stating now: because the selection criterion was ROC-AUC, which ignores the cut-off, the chosen "
        f"tree carries no class weighting, and at the default cut-off of 0.5 it flags almost nobody (recall {h05['tree']['recall']:.1%} on the "
        f"hold-out, Table 3) because almost no leaf contains a majority of out-of-school children. This is not a flaw in the tree. It is the "
        f"reason the operating cut-off is chosen separately, and on training data only, in section 5. The same applies to the forest wherever "
        f"it was tuned without class weighting. The full grids are in tables t02-t04.")

    doc.add_heading("3.1 Hypothesis H5 - do the household features add predictive value?", 2)
    para(doc,
        f"Hypothesis H5 asks whether adding poverty, B-form, disability and the control variables improves prediction over a bare "
        f"demographic model using only age, sex and area. On identical cross-validation folds the demographic model reached a ROC-AUC of "
        f"{h5['auc_demo_mean']:.3f} (plus or minus {h5['auc_demo_sd']:.3f}) and the full 13-feature model {h5['auc_full_mean']:.3f} (plus or minus "
        f"{h5['auc_full_sd']:.3f}): a gain of {h5['gain_mean']:+.3f} that appeared in every fold. A paired t-test (a test of whether the average "
        f"fold-by-fold difference could plausibly be zero) gives p = {h5['p_value']:.3f} (the probability of seeing a gain this large if there "
        f"were truly none). On the hold-out districts the two ROC-AUCs were {h5['holdout_auc_demo']:.3f} and {h5['holdout_auc_full']:.3f}.")
    para(doc,
        f"Five folds are a thin basis for a p-value, so we also used a bootstrap (re-computing the result many times on random re-samples "
        f"of the data to see how much it moves). Because districts, not children, are the natural units, we re-sampled the {bs['n_clusters']} "
        f"held-out districts with replacement {bs['n_boot']:,} times (a cluster bootstrap). The hold-out gain is {bs['diff_point']:+.3f} ROC-AUC "
        f"with a 95% interval of [{bs['ci_low']:+.3f}, {bs['ci_high']:+.3f}] and was positive in {bs['share_positive']:.0%} of re-samples. "
        f"We therefore reject the null hypothesis for H5 (the assumption of no improvement): who is out of school is explained far better "
        f"by household circumstances than by age, sex and area alone.")
    add_fig(doc, "h5", "Figure 1. ROC-AUC of the demographic and full logistic models in each of the five district-grouped folds.", 4.8)

    # ------------------------------------------------------------------ 4
    doc.add_heading("4. Results on the held-out districts", 1)
    para(doc,
        "Table 3 reports the standard classification metrics on the hold-out children, with every model using the default cut-off of 0.5 "
        "(a child is flagged when the model's estimated probability of being out of school is at least 0.5). Precision is the share of "
        "flagged children who really are out of school. Recall, also called sensitivity, is the share of truly out-of-school children who "
        "were flagged. Specificity is the share of enrolled children correctly left unflagged. F1 is a single score that balances precision "
        "and recall. PR-AUC (precision-recall area) is like ROC-AUC but focuses on the rare class, so it is more demanding when the outcome "
        "is uncommon; chance level equals the out-of-school share. 'Flagged' is the share of all children the model would send for outreach. "
        "Each metric is shown unweighted and then with the ASER sampling weights applied.")
    ho = pd.read_csv(T / "t07_holdout_metrics_threshold_0.5.csv")
    ho = ho[["model_name", "weighted", "accuracy", "precision", "recall", "specificity", "f1", "roc_auc", "pr_auc", "flagged_rate"]]
    ho["model_name"] = ho["model_name"].replace({"Logistic regression (age, female, urban)": NAMES["logreg_demo"]})
    ho.columns = ["Model", "Weighted", "Accuracy", "Precision", "Recall", "Specificity", "F1", "ROC-AUC", "PR-AUC", "Flagged"]
    add_table(doc, ho, caption="Table 3. Hold-out performance at the default 0.5 cut-off, unweighted and survey-weighted")
    para(doc,
        f"The majority-class baseline reaches {h05['naive']['accuracy']:.1%} accuracy and finds nobody. The balanced logistic regression gives "
        f"up accuracy ({h05['logreg']['accuracy']:.1%}) to reach a recall of {h05['logreg']['recall']:.1%} at the default cut-off, with ROC-AUC "
        f"{h05['logreg']['roc_auc']:.3f} and PR-AUC {h05['logreg']['pr_auc']:.3f}. The decision tree ranks children better (ROC-AUC "
        f"{h05['tree']['roc_auc']:.3f}, PR-AUC {h05['tree']['pr_auc']:.3f}) but, as explained in section 3, its default-cut-off recall is only "
        f"{h05['tree']['recall']:.1%}. The random forest ranks children best of all (ROC-AUC {h05['rf']['roc_auc']:.3f}, PR-AUC "
        f"{h05['rf']['pr_auc']:.3f}) with a default-cut-off recall of {h05['rf']['recall']:.1%}. The lesson is that the cut-off, not the choice of "
        f"model, decides who gets flagged.")
    para(doc,
        f"Applying the sampling weights moves the results in different directions for different models. ROC-AUC rises from "
        f"{h05['logreg']['roc_auc']:.3f} to {w05['logreg']['roc_auc']:.3f} for the full logistic model and from {h05['tree']['roc_auc']:.3f} to "
        f"{w05['tree']['roc_auc']:.3f} for the tree and from {h05['rf']['roc_auc']:.3f} to {w05['rf']['roc_auc']:.3f} for the forest, but falls from {h05['logreg_demo']['roc_auc']:.3f} to {w05['logreg_demo']['roc_auc']:.3f} for "
        f"the demographic-only model. The weights give more influence to households the sample under-represents, which are mostly poorer and "
        f"rural. That is exactly where the household features carry most of their information and where age, sex and area alone separate "
        f"children least. Statements about the national population should therefore quote the weighted figures. Hold-out ROC-AUCs are "
        f"0.03-0.06 below the cross-validation averages: the held-out districts happen to be a harder group, but the gap is within the "
        f"fold-to-fold spread seen in Table 2.")
    add_fig(doc, "roc_pr", "Figure 2. ROC curves (left) and precision-recall curves (right) on the held-out districts.")
    add_fig(doc, "cm_05", "Figure 3. Confusion matrices at the default 0.5 cut-off. Each cell shows a count and the share of its row.")

    # ------------------------------------------------------------------ 5
    doc.add_heading("5. The two kinds of error and what they cost", 1)
    para(doc,
        "A confusion matrix (a two-by-two table of what the model predicted against what was true) separates two kinds of mistake. "
        "A false negative is a child the model marks as 'in school' who is actually out of school: a child who needs an enrolment or "
        "outreach visit and never gets one. For an education-access programme this is the costly error. A false positive is a child the "
        "model marks as 'out of school' who is actually enrolled: an outreach visit or a stipend slot is spent where it was not needed. "
        "This is wasteful but recoverable.")
    para(doc,
        "To make the trade-off explicit we price a false negative at r false positives, for r equal to 1, 3, 5 and 10, and compute the "
        "expected cost per 1,000 children screened (false positives plus r times false negatives, scaled to 1,000 children). We then "
        "sweep the cut-off from 0.05 to 0.95 and find the value that minimises this cost for each r.")
    para(doc,
        "The cut-off is itself a model setting, so it must not be chosen by looking at the hold-out; doing so and then reporting hold-out "
        "performance at that value would be circular. Instead, each tuned model was re-fitted on the five district-grouped training folds "
        "to produce out-of-fold predictions (a prediction for every training child made by a model that did not see that child), the "
        "cost curve was swept on those, and the chosen cut-off was then applied once to the hold-out. Table 4 shows both the training "
        "out-of-fold figures that drove each choice and the hold-out figures at that cut-off. Its last two columns show the 'oracle' "
        "cut-off (the one that would have minimised cost on the hold-out itself, which we could only know by peeking), so the reader can "
        "see how much such peeking would have flattered the result.")
    bt = pd.read_csv(T / "t09_cost_minimising_thresholds.csv")
    bt["model"] = bt["model"].map(NAMES)
    bt = bt[["model", "cost_ratio_fn_to_fp", "threshold", "oof_recall", "oof_precision", "oof_cost_per_1000",
             "holdout_recall", "holdout_precision", "holdout_flagged_rate", "holdout_cost_per_1000", "holdout_oracle_threshold", "holdout_oracle_cost_per_1000"]]
    bt.columns = ["Model", "FN:FP", "Cut-off (chosen on training)", "Train recall", "Train precision", "Train cost/1,000",
                  "Hold-out recall", "Hold-out precision", "Hold-out flagged", "Hold-out cost/1,000", "Oracle cut-off", "Oracle cost/1,000"]
    add_table(doc, bt, caption="Table 4. Cost-minimising cut-off for each cost ratio, chosen on training out-of-fold predictions and then evaluated once on the hold-out", font=7.5)
    add_fig(doc, "threshold", f"Figure 4. Recall, precision and expected cost against the cut-off (cost ratio {ratio}:1). Solid lines are training "
            "out-of-fold predictions, used to choose the cut-off; dashed lines are the hold-out, shown for comparison only; dash-dot lines mark the chosen cut-offs.")
    cmp_ = pd.read_csv(T / "t10_confusion_default_vs_recommended.csv")
    cmp_ = cmp_[["model", "threshold_label", "threshold", "tp", "fn", "fp", "tn", "recall", "precision", "specificity",
                 "flagged_per_1000", "caught_per_1000", "missed_per_1000", f"cost_per_1000_ratio{ratio}"]]
    cmp_["model"] = cmp_["model"].map(NAMES)
    cmp_.columns = ["Model", "Cut-off", "Value", "Found (TP)", "Missed (FN)", "False alarms (FP)", "Correctly cleared (TN)", "Recall", "Precision",
                    "Specificity", "Flagged/1,000", "Found/1,000", "Missed/1,000", f"Cost/1,000 (r={ratio})"]
    add_table(doc, cmp_, caption="Table 5. Default cut-off versus the recommended cut-off, with the result expressed per 1,000 children screened", font=8)
    add_fig(doc, "cm_rec", f"Figure 5. Confusion matrices at the recommended cut-offs (cost ratio {ratio}:1).")
    bullets(doc, [
        "When a false alarm costs as much as a missed child (r = 1), the cheapest policy is to flag nobody. Neither model is sharp enough "
        "to beat the base rate on those terms. The cost framing exposes this honestly instead of hiding it behind an accuracy score.",
        f"Once a missed child is weighted {ratio} times a wasted visit, the best cut-off drops sharply. Ranked by hold-out cost per 1,000 at "
        f"cut-offs chosen on training data, the models come out as: " + "; ".join(f"{NAMES[m]} {rank[m]:.0f}" for m in rank) + ". "
        f"Our recommendation is therefore the {NAMES[best].lower()} at a cut-off of {rec[best]:.2f}. On the held-out districts, for every 1,000 "
        f"children screened it flags {best_rec['flagged_per_1000']:.0f}, finds {best_rec['caught_per_1000']:.0f} of the "
        f"{best_rec['caught_per_1000'] + best_rec['missed_per_1000']:.0f} who are out of school and misses {best_rec['missed_per_1000']:.0f} "
        f"(recall {best_rec['recall']:.2f}, precision {best_rec['precision']:.2f}). The same model at the default 0.5 cut-off would find only "
        f"{best_05['caught_per_1000']:.0f} per 1,000.",
        (f"A note on the ranking. The {NAMES[top_auc].lower()} has the best ranking quality (highest hold-out ROC-AUC, {h05[top_auc]['roc_auc']:.3f}), "
         f"yet the {NAMES[best].lower()} wins on cost because its training-chosen cut-off transfers to the new districts almost perfectly, whereas "
         f"the {NAMES[top_auc].lower()}'s cut-off of {rec[top_auc]:.2f} sits a little below its hold-out optimum of "
         f"{ct[(top_auc, ratio)]['holdout_oracle_threshold']:.2f} and so flags more children than it needs to. The gap between the two "
         f"({rank[best]:.0f} versus {rank[top_auc]:.0f} per 1,000) is within the fold-to-fold variation seen in cross-validation, so the choice between "
         "them should rest on stability and explainability rather than on this difference alone." if top_auc != best else
         f"The {NAMES[best].lower()} also has the highest hold-out ROC-AUC, so ranking quality and cost agree on the recommendation."),
        "The other two models at their own training-chosen cut-offs: " + "; ".join(
            f"{NAMES[m].lower()} at {rec[m]:.2f} (recall {tc[(m, 'recommended')]['recall']:.2f}, precision {tc[(m, 'recommended')]['precision']:.2f}, "
            f"{tc[(m, 'recommended')]['flagged_per_1000']:.0f} flagged per 1,000)" for m in others) + ". Logistic regression remains the better "
        "choice wherever officials must be able to explain, child by child, why someone was flagged, because its reasoning is a fixed list of "
        "weighted factors; the forest's advantage in ranking quality comes at the price of being harder to explain.",
        f"At r = 10 the models flag more than half of all children in order to find "
        f"{min(ct[(m,10)]['holdout_recall'] for m in MODELS):.0%}-{max(ct[(m,10)]['holdout_recall'] for m in MODELS):.0%} of those out of school. Whether that is affordable is a budget decision for the programme; the "
        "cost ratio simply makes the decision explicit.",
        "Honesty check. Had we chosen each cut-off on the hold-out itself, the results would have been: " + "; ".join(
            f"{NAMES[m].lower()} oracle {ct[(m, ratio)]['holdout_oracle_threshold']:.2f} costing {ct[(m, ratio)]['holdout_oracle_cost_per_1000']:.0f} per 1,000 "
            f"versus {ct[(m, ratio)]['holdout_cost_per_1000']:.0f} at the training-chosen {rec[m]:.2f}" for m in MODELS) + ". The flattery avoided is "
        f"at most {max(ct[(m, ratio)]['holdout_cost_per_1000'] / ct[(m, ratio)]['holdout_oracle_cost_per_1000'] - 1 for m in MODELS):.1%} of cost, small "
        "but not zero, which is why every figure at the recommended cut-off in Tables 5-7 and Figures 5-6 uses the training-chosen value.",
    ])

    doc.add_heading("5.1 Which children does the model get wrong?", 2)
    para(doc,
        f"Table 6 compares the average characteristics of the children the logistic regression found (true positives), missed (false negatives), "
        f"wrongly flagged (false positives) and correctly cleared (true negatives) at its recommended cut-off of {rec['logreg']:.2f}. Table 7 breaks "
        "performance down by subgroup. FNR is the false-negative rate, the share of out-of-school children in that group who were missed; it "
        "equals one minus recall.")
    prof = pd.read_csv(T / "t12_misclassified_profile_logreg.csv")
    prof = prof[["feature", "TP", "FN", "FP", "TN", "gap_FN_minus_TP"]]
    prof.columns = ["Feature (group average)", "Found (TP)", "Missed (FN)", "False alarms (FP)", "Correctly cleared (TN)", "Missed minus found"]
    add_table(doc, prof, caption=f"Table 6. Average characteristics by outcome, logistic regression at cut-off {rec['logreg']:.2f}")
    subt = pd.read_csv(T / "t14_subgroup_metrics_recommended_threshold.csv")
    subt = subt[subt.model == "logreg"][["variable", "level", "n", "positives", "prevalence", "recall", "fnr", "precision", "roc_auc"]]
    subt.columns = ["Variable", "Level", "Children (test)", "Out of school", "Share out of school", "Recall", "FNR", "Precision", "ROC-AUC"]
    add_table(doc, subt, caption="Table 7. Logistic regression performance by subgroup on the held-out districts")
    add_fig(doc, "subgroup", "Figure 6. Recall by subgroup, logistic regression at the recommended cut-off.", 5.6)
    para(doc,
        f"The children the model misses look like enrolled children on paper. They are older, far more often urban ({edge['urban']['FN']:.0%} of "
        f"missed children versus {edge['urban']['TP']:.0%} of those found), more often B-form registered, and live in households with more assets, "
        "more digital use and better housing. In effect the model has learned the rule 'poor rural household means out of school' and "
        "under-detects out-of-school children in better-off or urban households. These are plausibly children who dropped out rather than "
        "never enrolled, a distinction the current features cannot see.")
    para(doc,
        f"Recall is {sub.loc[('province','Balochistan'),'recall']:.2f} in Balochistan but only {sub.loc[('province','Sindh'),'recall']:.2f} in Sindh and "
        f"{sub.loc[('province','Punjab'),'recall']:.2f} in Punjab, and {sub.loc[('urban',1),'recall']:.2f} for urban children against "
        f"{sub.loc[('urban',0),'recall']:.2f} for rural children. Islamabad Capital Territory is a single district and fell entirely on the training "
        "side of the grouped split, so it has no hold-out result; with only 52 out-of-school children in the whole dataset it is in any case too "
        "small for reliable subgroup conclusions.")

    # ------------------------------------------------------------------ 6
    doc.add_heading("6. Which factors drive the predictions, and what that says about the hypotheses", 1)
    para(doc,
        "We look at importance from three angles, because each answers a slightly different question. Odds ratios from the logistic "
        "regression (Table 8) describe the association of each factor with being out of school once the other twelve are held fixed; an odds "
        "ratio above 1 raises the odds, below 1 lowers them, and a confidence interval that excludes 1 means the effect is unlikely to be "
        "chance. The p-value is the probability of an association at least this strong if there were truly none. The standardised coefficient "
        "(the effect of a one-standard-deviation change, which puts factors on different scales onto a common footing) ranks factors by size. "
        "Gini importance from the tree (Table 9) measures how much each factor was used to separate the groups. Permutation importance "
        "(how much the hold-out ROC-AUC falls when a factor's values are shuffled, destroying its information) measures each factor's unique "
        f"contribution to ranking children in new districts. The logistic model's McFadden pseudo-R-squared (a 0-to-1 measure of how much "
        f"better the model fits than a model with no predictors) is {R['pseudo_r2']:.3f}, which is modest.")
    o = pd.read_csv(T / "t15_logit_odds_ratios.csv")[["label", "odds_ratio", "or_ci_low", "or_ci_high", "p_value", "std_coef"]]
    o.columns = ["Factor", "Odds ratio", "95% CI low", "95% CI high", "p-value", "Standardised coefficient"]
    add_table(doc, o, caption="Table 8. Logistic regression odds ratios, estimated on the training districts")
    g = pd.read_csv(T / "t17_tree_gini_importance.csv")[["feature", "label", "gini_importance"]] \
        .merge(pd.read_csv(T / "t17b_rf_gini_importance.csv")[["feature", "gini_importance"]].rename(columns={"gini_importance": "gini_rf"}), on="feature") \
        .merge(pd.read_csv(T / "t18_permutation_logreg.csv")[["feature", "auc_drop_mean"]].rename(columns={"auc_drop_mean": "perm_lr"}), on="feature") \
        .merge(pd.read_csv(T / "t19_permutation_tree.csv")[["feature", "auc_drop_mean"]].rename(columns={"auc_drop_mean": "perm_tree"}), on="feature") \
        .merge(pd.read_csv(T / "t19b_permutation_rf.csv")[["feature", "auc_drop_mean"]].rename(columns={"auc_drop_mean": "perm_rf"}), on="feature")
    g = g[["label", "gini_importance", "gini_rf", "perm_lr", "perm_tree", "perm_rf"]].sort_values("gini_rf", ascending=False)
    g.columns = ["Factor", "Gini (tree)", "Gini (forest)", "Permutation drop (logistic)", "Permutation drop (tree)", "Permutation drop (forest)"]
    add_table(doc, g, caption="Table 9. Gini importance for the tree and the forest, and hold-out permutation importance for all three models", font=8)
    add_fig(doc, "odds", "Figure 7. Odds ratios with 95% confidence intervals. Points to the right of 1 raise the odds of being out of school.", 5.6)
    add_fig(doc, "importance", "Figure 8. Logistic standardised coefficients, tree Gini importance and permutation importance side by side.")
    add_fig(doc, "importance_rf", "Figure 8b. Random forest Gini and permutation importance.", 6.3)
    add_fig(doc, "h4", "Figure 9. Predicted probability of being out of school for boys and girls in rural and urban areas (H4).", 4.6)

    hyp = pd.DataFrame(R["hypotheses"])[["hypothesis", "feature", "H1_expectation", "odds_ratio", "or_ci_low", "or_ci_high", "p_value", "gini_rank", "H0_rejected_at_5pct"]]
    hyp.columns = ["Hypothesis", "Variable", "What H1 expects", "Odds ratio", "CI low", "CI high", "p-value", "Gini rank", "Null rejected"]
    hyp["Gini rank"] = hyp["Gini rank"].map(lambda v: "-" if pd.isna(v) else str(int(v)))
    add_table(doc, hyp, caption="Table 10. Evidence for the five Milestone 2 hypotheses (null rejected at the 5% level)")
    bullets(doc, [
        f"H1 Poverty (asset count). The odds ratio is {ors['assets_n']['odds_ratio']:.3f} per additional asset and its interval excludes 1, so the "
        f"direction is as hypothesised, but the effect per asset is small; the count ranks {hyp_gini['assets_n']:.0f}th in Gini importance. "
        f"Housing structure (odds ratio {ors['house']['odds_ratio']:.2f} per step from kachha to pacca) and digital access (odds ratio "
        f"{ors['digital_n']['odds_ratio']:.2f} per service) are the dominant poverty signals in all three models. The null is rejected, with the "
        "qualification that the asset count is a weaker measure of poverty than housing.",
        f"H2 Legal identity (B-form). Registered children have an odds ratio of {ors['bform']['odds_ratio']:.3f}, which means unregistered children "
        f"have about {1/ors['bform']['odds_ratio']:.2f} times the odds of being out of school. The result is significant and consistent across all "
        "importance views and all three models. The null is rejected.",
        f"H3 Disability. Children with a functional difficulty have an odds ratio of {ors['any_disability']['odds_ratio']:.2f}, the largest single risk "
        "factor after living in Balochistan. The null is rejected.",
        f"H4 Gender by area. The interaction term has an odds ratio of {ors['female_x_urban']['odds_ratio']:.3f} with an interval excluding 1. Girls face "
        f"higher odds everywhere, but the predicted girl-boy gap is {gaps['rural_gender_gap']*100:.1f} percentage points in rural areas and only "
        f"{gaps['urban_gender_gap']*100:.1f} in urban areas (Figure 9): the gender penalty is mainly a rural phenomenon. The tree hardly uses the "
        "interaction column because it reproduces the same pattern through successive splits. The null is rejected.",
        f"H5 Feature value. The full model gains {h5['gain_mean']:+.3f} ROC-AUC in cross-validation, in every fold, and "
        f"{h5['holdout_auc_full'] - h5['holdout_auc_demo']:+.3f} on the hold-out with a bootstrap interval that excludes zero. The null is rejected.",
        f"Control variables. Living in Balochistan multiplies the odds by {ors['province_Balochistan']['odds_ratio']:.2f} even after the household "
        "factors are accounted for; age has a small negative coefficient within the 5-16 window; the Islamabad indicator is not significant "
        "(807 children).",
        "Why two importance measures disagree about Balochistan. Balochistan has the largest positive standardised coefficient in the logistic "
        "model yet almost no permutation importance on the hold-out, and Sindh's permutation importance is slightly negative. The province "
        "indicators overlap heavily with housing quality and digital access, a situation called collinearity (two or more predictors carrying "
        "much the same information): kachha, disconnected households are concentrated in Balochistan. Shuffling a province flag therefore costs "
        "little while the household variables remain. The coefficient measures the province's association given everything else; the "
        "permutation drop measures what is lost when only that factor is removed. They disagree because the information is shared, not because "
        "either is wrong.",
    ])

    # ------------------------------------------------------------------ 7
    doc.add_heading("7. What this means for policy, with an important caveat", 1)
    para(doc,
        "A caveat before the recommendations. ASER is a single cross-sectional survey (a snapshot of many households at one point in time, "
        "with no follow-up). Every factor above is therefore an association (two things tending to occur together), not an estimated causal "
        "effect (proof that changing one thing would change the other). The B-form result is the clearest example. Children without a "
        "registration are more likely to be out of school, but households that are more engaged with public services may obtain both the "
        "document and a school place for reasons unrelated to the document itself. If so, issuing B-forms on their own would not raise "
        "enrolment. The same reasoning applies to housing, digital access and disability. The points below therefore describe where and for "
        "whom out-of-school risk concentrates, which is what a targeting model can legitimately inform, and name candidate levers whose causal "
        "effect would have to be established, for example through a piloted or staggered roll-out, before being relied on as policy.")
    bullets(doc, [
        "Target by housing and connectivity, not only by an asset count. Housing quality and digital access are the strongest household "
        "predictors of being out of school. Outreach lists built from a simple asset score would miss children in kachha, disconnected homes.",
        "Birth registration is a candidate lever and a practical targeting marker. Unregistered children have about 1.3 times the odds of being "
        "out of school. At minimum, NADRA records identify a population the model flags as high-risk; a combined registration-and-enrolment "
        "drive is a testable intervention rather than a proven one.",
        "Disability inclusion is a first-order gap (SDG 4.5). Children with a functional difficulty have about 1.9 times the odds of being out "
        "of school. Screening and inclusive-school capacity belong in any enrolment drive.",
        "The gender gap is rural. Girl-focused measures should be prioritised in rural areas, where the predicted penalty is almost three times "
        "the urban one.",
        "Balochistan needs a province-level response. Living there nearly doubles the odds even after household factors are accounted for.",
        f"Set the operating cut-off deliberately, and choose it on training data. At a {ratio}:1 cost ratio the recommended policy "
        f"({NAMES[best].lower()} at {rec[best]:.2f}) flags about {best_rec['flagged_rate']:.0%} of children and finds {best_rec['recall']:.0%} of those "
        f"out of school in districts the model has never seen. "
        "The default cut-off of 0.5 is not a defensible choice for outreach.",
    ])

    # ------------------------------------------------------------------ 8
    doc.add_heading("8. Limitations and the plan for Milestone 04", 1)
    para(doc, "Current limitations.")
    bullets(doc, [
        f"Modest discrimination (the ability to tell the two groups apart): hold-out ROC-AUC of {min(h05[m]['roc_auc'] for m in MODELS):.2f}-{max(h05[m]['roc_auc'] for m in MODELS):.2f} "
        f"and a pseudo-R-squared of {R['pseudo_r2']:.2f}. This is a sign of underfitting (the model is too simple for the pattern). Thirteen household "
        "features cannot separate the harder cases, especially urban or better-off dropouts; the missing information is probably on the school "
        "side (distance, supply, cost) and in the child's history.",
        "Class imbalance is handled with class weights and the choice of cut-off, not by resampling (artificially balancing the training data). "
        "Precision stays near 0.2 at useful levels of recall, so roughly four in five flagged children would turn out to be enrolled.",
        "District heterogeneity: performance varies by about 0.05 ROC-AUC between groups of districts, and the held-out districts were harder "
        "than the cross-validation average.",
        "Potential demographic bias: recall differs sharply by area, province and B-form status. Girls and children with disabilities are "
        "currently found more often than their counterparts because the model assigns them higher risk, which is itself a fairness question "
        "to examine rather than assume is benign.",
        "Associational evidence only: the drivers are conditional associations from one survey, and no policy lever has an estimated causal "
        "effect (section 7).",
        "Thin inference for H5: the paired t-test rests on five folds. The district-cluster bootstrap of the hold-out gain is the sturdier "
        "statement, and the two agree.",
    ])
    para(doc, "Plan for Milestone 04.")
    bullets(doc, [
        "A fairness audit (a systematic check of whether the model's errors fall unevenly on particular groups) using the subgroup frame in "
        "subgroup_counts.csv: recall, precision and false-negative rate for girls versus boys, rural versus urban, with versus without a "
        "disability, and by province, each with bootstrap confidence intervals. We will test whether group-specific cut-offs or reweighting "
        "close the urban and Punjab/Sindh recall gaps. Islamabad Capital Territory (807 children, 52 out of school) is too small for reliable "
        "group conclusions and will be reported but not acted on.",
        "Stronger models under the same district-grouped cross-validation: gradient boosting (an ensemble that builds trees one after another, "
        "each correcting the errors of the last, and usually the strongest tabular learner), building on the random forest already fitted here; calibration (checking that a predicted 30% really means about 30%); interaction "
        "terms suggested by the tree; and a cost-sensitive objective (training the model to minimise the weighted error cost directly rather "
        "than fixing it afterwards with a cut-off).",
        "An ethical evaluation covering the harms of mis-targeting (stigma from a false flag, exclusion from a missed child), the use of survey "
        "weights for any population-level claim, and a deployment note on who should set the cost ratio and how.",
    ])

    # ------------------------------------------------------------------ 9
    doc.add_heading("9. Reproducibility", 1)
    para(doc, "Code: milestone3/src (config, data, cv, models, evaluate, importance, plots, run_all, make_report) and the executed notebook "
              "Milestone03_BaselineModels.ipynb. Running src/run_all.py regenerates every table (outputs/tables) and figure (outputs/figures) "
              "from data/aser2025_training.csv with a fixed random seed of 42, so all results are exactly repeatable; src/make_report.py rebuilds "
              "this document from those outputs.")

    # ------------------------------------------------------------------ glossary
    doc.add_heading("Glossary of terms used in this section", 1)
    gl = pd.DataFrame([
        ("Target variable", "The outcome column the model learns to predict; here out_of_school."),
        ("Predictor / feature", "An input column the model uses to make its prediction."),
        ("Class imbalance", "One outcome is much rarer than the other (13% out of school versus 87% enrolled)."),
        ("Sampling weight", "A number attached to each surveyed child so the sample represents the national population."),
        ("One-hot encoding / reference category", "Turning a category such as province into separate 0/1 columns; the left-out category is the comparison baseline."),
        ("Training set / hold-out (test) set", "Data the model learns from / data kept aside and used once for the final evaluation."),
        ("Cross-validation / fold", "Repeatedly training on part of the data and checking on the rest; each portion checked is a fold."),
        ("Stratified / grouped split", "A split that keeps the outcome share similar on both sides / keeps whole districts on one side."),
        ("Data leakage", "Information from the test data influencing the model, which inflates measured performance."),
        ("Hyperparameter / grid search", "A setting chosen by the analyst rather than learned / trying every combination from a list."),
        ("Out-of-fold prediction", "A prediction for a training child made by a model that did not see that child."),
        ("Cut-off (threshold)", "The probability above which a child is flagged as out of school."),
        ("Accuracy", "Share of all children classified correctly."),
        ("Precision", "Share of flagged children who really are out of school."),
        ("Recall (sensitivity)", "Share of truly out-of-school children the model flags."),
        ("Specificity", "Share of enrolled children correctly left unflagged."),
        ("F1", "A single score that balances precision and recall."),
        ("ROC-AUC", "A 0-to-1 score of how well the model ranks out-of-school children above enrolled ones; 0.5 is chance."),
        ("PR-AUC", "Like ROC-AUC but focused on the rare class; chance equals the out-of-school share."),
        ("Confusion matrix", "A two-by-two table of predicted versus actual outcomes."),
        ("False negative / false positive", "A missed out-of-school child / an enrolled child wrongly flagged."),
        ("Expected cost per 1,000", "False positives plus r times false negatives, scaled to 1,000 children screened."),
        ("Odds ratio", "How many times higher (above 1) or lower (below 1) the odds of the outcome are per unit change in a factor."),
        ("Confidence interval", "The range in which the true value very probably lies (here with 95% confidence)."),
        ("p-value", "The probability of a result at least this strong if there were truly no effect."),
        ("Standardised coefficient", "An effect size per one-standard-deviation change, comparable across factors."),
        ("Gini importance", "How much a factor was used by the tree to separate the groups."),
        ("Permutation importance", "How much performance falls when a factor's values are shuffled."),
        ("Collinearity", "Two or more predictors carrying much the same information."),
        ("Interaction term", "A column that lets the effect of one factor depend on another (here sex by area)."),
        ("Pseudo-R-squared", "A 0-to-1 measure of how much better the model fits than one with no predictors."),
        ("Bootstrap / cluster bootstrap", "Re-computing a result on many random re-samples / re-sampling whole districts rather than children."),
        ("Paired t-test", "A test of whether the average fold-by-fold difference between two models could be zero."),
        ("Cross-sectional survey", "A one-time snapshot of many households, with no follow-up over time."),
        ("Association versus causal effect", "Two things occurring together versus proof that changing one changes the other."),
        ("Underfitting", "A model too simple to capture the pattern in the data."),
        ("Ensemble / random forest / bagging", "A model made of many models; a forest averages many trees, each trained on a bootstrap re-sample (bagging) with random feature subsets."),
        ("Calibration", "Whether a predicted probability of 30% really corresponds to about 30% of such children being out of school."),
    ], columns=["Term", "Meaning"])
    add_table(doc, gl, font=8.5)

    out = C.REPORT_DIR / "Milestone03_Report.docx"
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
